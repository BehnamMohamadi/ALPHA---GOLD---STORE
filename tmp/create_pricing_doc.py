from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

out=Path('output/documents');out.mkdir(parents=True,exist_ok=True)
d=Document(); sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Arial';s.font.size=Pt(11 if name=='Normal' else 22 if name=='Title' else 14)
 s.font.color.rgb=RGBColor(0,0,0);s.element.get_or_add_rPr().append(OxmlElement('w:rtl'))
 s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.25
def rtl(p):
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
 for r in p.runs:
  pr=r._r.get_or_add_rPr();pr.append(OxmlElement('w:rtl'))
  fonts=OxmlElement('w:rFonts');fonts.set(qn('w:cs'),'Arial');pr.append(fonts)
 return p
def p(t,style=None):return rtl(d.add_paragraph(t,style))
def h(t):p(t,'Heading 1')
def table(rows,widths):
 t=d.add_table(rows=0, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 pr=t._tbl.tblPr;bi=OxmlElement('w:bidiVisual');pr.append(bi)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 pr.append(borders)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,txt in enumerate(row):
   c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=txt
   cp=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF1');cp.append(sh)
   for para in c.paragraphs:
    rtl(para);para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.15
    for r in para.runs:r.font.size=Pt(10);r.bold=i==0
  if i==0:t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p('')
p('راهنمای قیمت گذاری محصولات آلفا','Title')
p('این راهنما روش محاسبه قیمت استاندارد و سفارشی در نسخه فعلی فروشگاه آلفا و شیوه تنظیم آن در پنل مدیریت را توضیح می‌دهد. قیمت‌گذاری سفارشی، سود، مالیات و اعمال اجرت را برای یک محصول مشخص تنظیم می‌کند؛ قیمت نهایی همچنان با نرخ طلای ثبت‌شده در پنل محاسبه می‌شود و ثابت نیست.')
h('تفاوت قیمت گذاری استاندارد و سفارشی')
table([
 ['تنظیم','استاندارد','سفارشی'],
 ['نرخ طلا','نرخ عیار محصول در پنل قیمت طلا','همان نرخ'],
 ['درصد سود','مقدار عمومی فروشگاه','مقدار اختصاصی محصول'],
 ['درصد مالیات','مقدار عمومی فروشگاه','مقدار اختصاصی محصول'],
 ['اجرت','نوع و مقدار اجرت محصول','همان اجرت با امکان خاموش کردن'],
 ['متعلقات','مبلغ ثبت‌شده برای محصول','همان مبلغ'],
],[1.3,2.75,2.75])
p('در حالت سفارشی، خالی گذاشتن درصد سود یا مالیات یعنی استفاده از مقدار عمومی فروشگاه؛ واردکردن صفر یعنی صفر درصد. در حالت استاندارد، درصدهای سفارشی و کلید خاموش‌کردن اجرت اثری ندارند.')
h('محاسبه قیمت مرحله به مرحله')
p('۱  ارزش طلای محصول: وزن طلا به گرم × نرخ هر گرم طلای همان عیار. برای محصول ۱۸ عیار، نرخ ۱۸ عیار استفاده می‌شود.')
p('۲  اجرت ساخت: در حالت درصدی، ارزش طلا × درصد اجرت ÷ ۱۰۰ محاسبه می‌شود. در حالت ثابت، همان مبلغ واردشده به تومان برای هر عدد محصول منظور می‌شود، نه برای هر گرم. در قیمت‌گذاری سفارشی، خاموش‌بودن «اعمال اجرت» مبلغ اجرت را صفر می‌کند.')
p('۳  سود فروشگاه: مجموع ارزش طلا و اجرت × درصد سود ÷ ۱۰۰.')
p('۴  مالیات در فرمول فعلی سایت: مجموع اجرت و سود × درصد مالیات ÷ ۱۰۰.')
p('۵  قیمت نهایی هر عدد محصول: ارزش طلا + اجرت + سود + مالیات + قیمت متعلقات.')
p('در کد فعلی، مبلغ متعلقات در پایان اضافه می‌شود و وارد محاسبه سود و مالیات نمی‌شود. هر بخش به نزدیک‌ترین تومان گرد می‌شود.')
d.add_page_break()
h('مثال قیمت گذاری با سود سفارشی ۲۲ درصد')
p('مقادیر زیر فقط برای توضیح محاسبه هستند و بیانگر نرخ جاری طلا یا نرخ قانونی مالیات نیستند.')
table([['ورودی','مقدار'],['وزن طلای ۱۸ عیار','۲ گرم'],['نرخ هر گرم','۷٬۴۰۰٬۰۰۰ تومان'],['اجرت','۱۰٪'],['سود سفارشی','۲۲٪'],['مالیات سفارشی','۱۰٪'],['متعلقات','۲۰۰٬۰۰۰ تومان']],[3.4,3.4])
h('نتیجه محاسبه')
table([['بخش','مبلغ به تومان'],['ارزش طلا','۱۴٬۸۰۰٬۰۰۰'],['اجرت','۱٬۴۸۰٬۰۰۰'],['سود ۲۲٪ از ۱۶٬۲۸۰٬۰۰۰','۳٬۵۸۱٬۶۰۰'],['مالیات ۱۰٪ از ۵٬۰۶۱٬۶۰۰','۵۰۶٬۱۶۰'],['متعلقات','۲۰۰٬۰۰۰'],['قیمت نهایی','۲۰٬۵۶۷٬۷۶۰']],[3.4,3.4])
p('بنابراین عدد ۲۲ در فیلد «درصد سود سفارشی» فقط درصد سود محصول است. اجرت و مالیات جداگانه و با فرمول‌های بالا محاسبه می‌شوند.')
h('تنظیم قیمت در پنل مدیریت')
p('۱  وزن طلا و عیار محصول را وارد کنید.\n۲  نوع اجرت را درصدی یا ثابت انتخاب و مقدار آن را مشخص کنید.\n۳  قیمت متعلقات را به تومان وارد کنید؛ اگر محصول متعلقات ندارد، صفر بگذارید.\n۴  نوع قیمت‌گذاری را روی «سفارشی» قرار دهید.\n۵  درصد سود و مالیات دلخواه را وارد کنید؛ برای استفاده از درصد عمومی، فیلد مربوط را خالی بگذارید.\n۶  وضعیت «اعمال اجرت» را مشخص و محصول را ذخیره کنید.')
p('با تغییر نرخ طلا در پنل، قیمت محاسبه‌شده محصول نیز تغییر می‌کند. همه مبالغ این راهنما به تومان هستند و قیمت نهایی توضیح‌داده‌شده، قیمت هر عدد محصول است.')
d.core_properties.title='راهنمای قیمت گذاری محصولات آلفا';d.core_properties.author='ALPHA'
d.save(out/'alpha-pricing-guide.docx')
print(out/'alpha-pricing-guide.docx')
